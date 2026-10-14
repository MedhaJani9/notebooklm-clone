import io

import pytest

from nbclone import ingestion
from nbclone.retrieval import NotebookIndex
from tests.conftest import FakeEmbedder


def make_pdf(pages):
    from reportlab.pdfgen import canvas
    buf = io.BytesIO()
    c = canvas.Canvas(buf)
    for t in pages:
        c.drawString(72, 750, t)
        c.showPage()
    c.save()
    return buf.getvalue()


def make_pptx(slides):
    from pptx import Presentation
    prs = Presentation()
    for t in slides:
        s = prs.slides.add_slide(prs.slide_layouts[1])
        s.shapes.title.text = t
        s.placeholders[1].text = "body of " + t
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def test_pdf_pages_tracked():
    segs = ingestion.extract_file("a.pdf", make_pdf(["Alpha page", "Beta page"]))
    assert [s["page"] for s in segs] == [1, 2] and "Beta" in segs[1]["text"]


def test_pptx_slides_tracked():
    segs = ingestion.extract_file("a.pptx", make_pptx(["Intro", "Details"]))
    assert segs[1]["label"] == "slide 2" and "Details" in segs[1]["text"]


def test_txt_and_chunking():
    segs = ingestion.extract_file("a.txt", ("word " * 1000).encode())
    chunks = ingestion.chunk_segments(segs, "sid", "a.txt")
    assert len(chunks) > 1 and all(len(c["text"]) <= 800 for c in chunks)
    assert chunks[0]["source"] == "a.txt" and chunks[0]["source_id"] == "sid"


def test_unsupported_and_empty_errors():
    with pytest.raises(ingestion.IngestionError):
        ingestion.extract_file("a.exe", b"x")
    with pytest.raises(ingestion.IngestionError):
        ingestion.chunk_segments(ingestion.extract_file("a.txt", b"   "), "s", "a.txt")
    with pytest.raises(ingestion.IngestionError):
        ingestion.extract_file("bad.pdf", b"not a pdf")


def test_url_extraction(monkeypatch):
    import requests

    class R:
        text = "<html><title>T</title><nav>menu</nav><body><p>Hello world content</p><script>x()</script></body></html>"
        headers = {"content-type": "text/html"}
        def raise_for_status(self): pass

    monkeypatch.setattr(requests, "get", lambda *a, **k: R())
    monkeypatch.setattr(ingestion, "_check_url", lambda u: None)
    title, segs = ingestion.extract_url("https://example.com")
    assert title == "T" and "Hello world content" in segs[0]["text"] and "menu" not in segs[0]["text"]


def test_url_ssrf_blocked():
    for bad in ("http://127.0.0.1/x", "http://localhost:8080", "file:///etc/passwd", "ftp://x.com"):
        with pytest.raises(ingestion.IngestionError):
            ingestion.extract_url(bad)


DOCS = ["Python lists are mutable ordered containers that support append and pop.",
        "Dictionaries map keys to values and are created with curly braces.",
        "The zylophrax protocol uses a unique handshake identifier XJ-9000."]


def build_index():
    emb = FakeEmbedder()
    chunks = [{"text": t, "source": f"s{i}", "page": None, "label": None, "embedding": e}
              for i, (t, e) in enumerate(zip(DOCS, emb.embed(DOCS)))]
    return NotebookIndex("abcdef123456", chunks, emb)


@pytest.mark.parametrize("mode", ["vector", "hybrid"])
def test_retrieval_finds_relevant_chunk(mode):
    res = build_index().search("what are python lists", k=2, mode=mode)
    assert res[0]["source"] == "s0" and res[0]["method"] == mode and "embedding" not in res[0]


def test_hybrid_catches_exact_keyword():
    assert build_index().search("XJ-9000", k=1, mode="hybrid")[0]["source"] == "s2"


def test_empty_index_and_bad_mode():
    assert NotebookIndex("abcdef000000", [], FakeEmbedder()).search("x") == []
    with pytest.raises(ValueError):
        build_index().search("x", mode="nope")
