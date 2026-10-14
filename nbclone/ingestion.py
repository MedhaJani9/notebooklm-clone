"""Text extraction + chunking for PDF, PPTX, TXT and single-page URLs.

Every extracted segment keeps its location (PDF page / PPTX slide) so each
chunk can later be cited precisely.
"""
import io
import ipaddress
import os
import socket
from urllib.parse import urlparse

from langchain_text_splitters import RecursiveCharacterTextSplitter

from . import config

SUPPORTED = {".pdf", ".pptx", ".txt"}


class IngestionError(Exception):
    """Raised with a user-friendly message when a source cannot be ingested."""


def extract_pdf(data: bytes) -> list:
    from pypdf import PdfReader
    try:
        reader = PdfReader(io.BytesIO(data))
        return [{"text": (p.extract_text() or "").strip(), "page": i + 1, "label": f"p. {i + 1}"}
                for i, p in enumerate(reader.pages)]
    except Exception as e:
        raise IngestionError(f"Could not read this PDF ({e.__class__.__name__}). Is it corrupted or password-protected?")


def extract_pptx(data: bytes) -> list:
    from pptx import Presentation
    try:
        prs = Presentation(io.BytesIO(data))
    except Exception as e:
        raise IngestionError(f"Could not read this PPTX ({e.__class__.__name__}).")
    segments = []
    for i, slide in enumerate(prs.slides, 1):
        parts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                parts.append(shape.text_frame.text)
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    parts.append(" | ".join(c.text for c in row.cells))
        if slide.has_notes_slide:
            parts.append("Notes: " + slide.notes_slide.notes_text_frame.text)
        segments.append({"text": "\n".join(p for p in parts if p.strip()), "page": i, "label": f"slide {i}"})
    return segments


def extract_txt(data: bytes) -> list:
    for enc in ("utf-8", "latin-1"):
        try:
            return [{"text": data.decode(enc).strip(), "page": None, "label": None}]
        except UnicodeDecodeError:
            continue
    raise IngestionError("Could not decode this text file.")


def _check_url(url: str) -> None:
    """Reject non-http(s) URLs and hosts resolving to private/loopback addresses (SSRF guard)."""
    p = urlparse(url)
    if p.scheme not in ("http", "https") or not p.hostname:
        raise IngestionError("Please enter a valid http(s) URL.")
    try:
        for info in socket.getaddrinfo(p.hostname, None):
            ip = ipaddress.ip_address(info[4][0])
            if ip.is_private or ip.is_loopback or ip.is_link_local or ip.is_reserved:
                raise IngestionError("That URL points to a private/internal address and was blocked.")
    except socket.gaierror:
        raise IngestionError("Could not resolve that URL's hostname.")


def extract_url(url: str) -> tuple:
    import requests
    from bs4 import BeautifulSoup

    url = url.strip()
    _check_url(url)
    try:
        resp = requests.get(url, timeout=20, headers={"User-Agent": "Mozilla/5.0 NotebookClone/1.0"})
        resp.raise_for_status()
    except Exception as e:
        raise IngestionError(f"Could not fetch that URL ({e.__class__.__name__}).")
    if "html" not in resp.headers.get("content-type", "html") and "text" not in resp.headers.get("content-type", ""):
        raise IngestionError("That URL is not a web page (HTML/text).")
    soup = BeautifulSoup(resp.text, "html.parser")
    for tag in soup(["script", "style", "nav", "footer", "header", "aside", "form", "noscript"]):
        tag.decompose()
    title = (soup.title.string or "").strip() if soup.title and soup.title.string else url
    text = "\n".join(line.strip() for line in soup.get_text("\n").splitlines() if line.strip())
    return title[:120], [{"text": text, "page": None, "label": None}]


def extract_file(filename: str, data: bytes) -> list:
    ext = os.path.splitext(filename)[1].lower()
    if ext not in SUPPORTED:
        raise IngestionError(f"Unsupported file type '{ext or filename}'. Use PDF, PPTX or TXT.")
    if len(data) > config.MAX_UPLOAD_MB * 1024 * 1024:
        raise IngestionError(f"File is larger than {config.MAX_UPLOAD_MB} MB.")
    return {".pdf": extract_pdf, ".pptx": extract_pptx, ".txt": extract_txt}[ext](data)


def chunk_segments(segments: list, source_id: str, source_name: str) -> list:
    """Split each segment separately so a chunk never spans two pages/slides."""
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=config.CHUNK_SIZE, chunk_overlap=config.CHUNK_OVERLAP,
        separators=["\n\n", "\n", ". ", " ", ""])
    chunks = []
    for seg in segments:
        if not seg["text"].strip():
            continue
        for piece in splitter.split_text(seg["text"]):
            chunks.append({"text": piece, "source_id": source_id, "source": source_name,
                           "page": seg["page"], "label": seg["label"], "chunk_index": len(chunks)})
    if not chunks:
        raise IngestionError("No text could be extracted (scanned/image-only documents are not supported).")
    return chunks
